"""Helpers that edit the user's own deck in place: clone their content slide (background frame, title, red rule,
body and footnote boxes), write text in their styling (Times New Roman, navy 0F1E42 / slate 1E293B / red B91C1C),
add tables in their header style, pictures and speaker notes."""
import copy
import re

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

NAVY, BODY, RED, GRAY, LIGHT = "0F1E42", "1E293B", "B91C1C", "64748B", "F1F5F9"
FONT = "Times New Roman"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def I(x):
    return int(round(x * 914400))


# ------------------------------------------------------------------ slides
def clone_slide(prs, src):
    """Append a copy of ``src`` (same layout, same shapes, image relationships re-created)."""
    dst = prs.slides.add_slide(src.slide_layout)
    for shp in list(dst.shapes):
        shp._element.getparent().remove(shp._element)
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype == RT.IMAGE:
            rid_map[rid] = dst.part.relate_to(rel.target_part, RT.IMAGE)
    tree = dst.shapes._spTree
    for el in src.shapes._spTree.iterchildren():
        tag = etree.QName(el).localname
        if tag in ("nvGrpSpPr", "grpSpPr", "extLst"):
            continue
        new = copy.deepcopy(el)
        for blip in new.iter(qn("a:blip")):
            old = blip.get(f"{{{R_NS}}}embed")
            if old in rid_map:
                blip.set(f"{{{R_NS}}}embed", rid_map[old])
        tree.append(new)
    return dst


def delete_slide(prs, slide):
    lst = prs.slides._sldIdLst
    for sldId in list(lst):
        if prs.part.related_part(sldId.rId) is slide.part:
            prs.part.drop_rel(sldId.rId)
            lst.remove(sldId)
            return
    raise ValueError("slide not found")


def move_slide_to_end(prs, slide):
    lst = prs.slides._sldIdLst
    for sldId in list(lst):
        if prs.part.related_part(sldId.rId) is slide.part:
            lst.remove(sldId)
            lst.append(sldId)
            return


def shape(slide, name):
    for s in slide.shapes:
        if s.name == name:
            return s
    raise KeyError(name)


def remove(shp):
    shp._element.getparent().remove(shp._element)


# ------------------------------------------------------------------ text
_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def _runs(p, text, size, color, bold=False, italic=False, lang="en-US"):
    parts = _INLINE.split(text)
    for part in parts:
        if not part:
            continue
        b, i, t = bold, italic, part
        if part.startswith("**") and part.endswith("**"):
            b, t = True, part[2:-2]
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            i, t = True, part[1:-1]
        r = etree.SubElement(p, qn("a:r"))
        rpr = etree.SubElement(r, qn("a:rPr"), lang=lang, sz=str(int(size * 100)), b="1" if b else "0", i="1" if i else "0", dirty="0")
        fill = etree.SubElement(rpr, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr"), val=color)
        etree.SubElement(rpr, qn("a:latin"), typeface=FONT)
        etree.SubElement(rpr, qn("a:cs"), typeface=FONT)
        etree.SubElement(r, qn("a:t")).text = t


def para(txBody, spec, lang="en-US"):
    """spec: (kind, text) or dict. kinds: h (heading 16 bold navy), b (bullet 14), s (sub-bullet 13), p (plain 14),
    n (note 12 italic grey), c (centred), big (stat number)."""
    if isinstance(spec, tuple):
        spec = {"k": spec[0], "t": spec[1]}
    k = spec.get("k", "p")
    size = spec.get("size", {"h": 16, "b": 14, "s": 13, "p": 14, "n": 12, "c": 14, "big": 26}[k])
    color = spec.get("color", {"h": NAVY, "big": RED, "n": GRAY}.get(k, BODY))
    bold = spec.get("bold", k in ("h", "big"))
    italic = spec.get("italic", k == "n")
    before = spec.get("before", {"h": 7, "b": 3, "s": 1, "p": 3, "n": 3, "c": 2, "big": 0}[k])
    p = etree.SubElement(txBody, qn("a:p"))
    attrs = {"algn": spec.get("align", "ctr" if k in ("c", "big") else "l")}
    if k == "b":
        attrs.update(marL=str(I(0.24)), indent=str(-I(0.2)))
    elif k == "s":
        attrs.update(marL=str(I(0.52)), indent=str(-I(0.18)))
    ppr = etree.SubElement(p, qn("a:pPr"), **attrs)
    sb = etree.SubElement(ppr, qn("a:spcBef"))
    etree.SubElement(sb, qn("a:spcPts"), val=str(int(before * 100)))
    if k in ("b", "s"):
        etree.SubElement(ppr, qn("a:buFont"), typeface="Arial")
        etree.SubElement(ppr, qn("a:buChar"), char="•" if k == "b" else "–")
    else:
        etree.SubElement(ppr, qn("a:buNone"))
    _runs(p, spec["t"], size, color, bold, italic, lang)
    return p


def set_text(shp, specs, lang="en-US", first_before0=True):
    txBody = shp.text_frame._txBody
    for p in txBody.findall(qn("a:p")):
        txBody.remove(p)
    for n, s in enumerate(specs):
        if isinstance(s, tuple):
            s = {"k": s[0], "t": s[1]}
        if n == 0 and first_before0:
            s = {**s, "before": 0}
        para(txBody, s, lang)


def place(shp, x, y, w, h=None):
    shp.left, shp.top, shp.width = I(x), I(y), I(w)
    if h is not None:
        shp.height = I(h)


def textbox_like(slide, proto, x, y, w, h, specs, lang="en-US"):
    """New text box that copies the body box's XML (insets, autofit) and gets new text."""
    el = copy.deepcopy(proto._element)
    slide.shapes._spTree.append(el)
    shp = slide.shapes[-1]
    shp.name = f"TextBox {len(slide.shapes) + 1}"
    place(shp, x, y, w, h)
    set_text(shp, specs, lang)
    return shp


def card(slide, x, y, w, h, big, label, fill=LIGHT, big_color=RED, big_size=24, label_size=12):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, I(x), I(y), I(w), I(h))
    shp.adjustments[0] = 0.08
    shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    shp.line.color.rgb = RGBColor.from_string("CBD5E1"); shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = I(0.08); tf.margin_top = tf.margin_bottom = I(0.05)
    tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_text(shp, [{"k": "big", "t": big, "color": big_color, "size": big_size}, {"k": "c", "t": label, "size": label_size, "color": BODY, "before": 3}])
    return shp


# ------------------------------------------------------------------ tables (user's style: navy header, white bold text)
def table(slide, x, y, widths, header, rows, size=12, row_h=0.27, highlight=(), align=None, bold_first=False, head_size=None):
    nr, nc = len(rows) + 1, len(header)
    gf = slide.shapes.add_table(nr, nc, I(x), I(y), I(sum(widths)), I(row_h * nr))
    gf.name = f"Table {len(slide.shapes)}"
    t = gf.table
    t.first_row = True; t.horz_banding = False
    for j, w in enumerate(widths):
        t.columns[j].width = I(w)
    for i in range(nr):
        t.rows[i].height = I(row_h)
    align = align or (["l"] + ["ctr"] * (nc - 1))
    for i in range(nr):
        for j in range(nc):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = I(0.06); c.margin_top = c.margin_bottom = I(0.025)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            if i == 0:
                fill, color, bold, txt = NAVY, "FFFFFF", True, header[j]
            else:
                hl = (i - 1) in highlight
                fill = "FEE2E2" if hl else ("FFFFFF" if (i - 1) % 2 == 0 else "F8FAFC")
                color, bold, txt = NAVY, hl or (bold_first and j == 0), rows[i - 1][j]
            c.fill.solid(); c.fill.fore_color.rgb = RGBColor.from_string(fill)
            tb = c.text_frame._txBody
            for p in tb.findall(qn("a:p")):
                tb.remove(p)
            para(tb, {"k": "p", "t": str(txt), "size": head_size if (i == 0 and head_size) else size, "color": color, "bold": bold,
                      "before": 0, "align": "ctr" if i == 0 else align[j]})
    return gf


def picture(slide, path, x, y, w=None, h=None):
    from PIL import Image
    iw, ih = Image.open(path).size
    if w is not None and h is not None:                 # fit inside the box, keep aspect, centre
        s = min(w / iw, h / ih)
        pw, ph = iw * s, ih * s
        return slide.shapes.add_picture(path, I(x + (w - pw) / 2), I(y + (h - ph) / 2), I(pw), I(ph))
    if w is not None:
        return slide.shapes.add_picture(path, I(x), I(y), I(w), I(w * ih / iw))
    return slide.shapes.add_picture(path, I(x), I(y), I(h * iw / ih), I(h))


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text
