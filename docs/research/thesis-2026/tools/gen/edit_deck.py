"""Edit the user's own deck in place (their background frames, colours, fonts and layout), correcting every number to
the measured runs and expanding the content.  usage: python edit_deck.py <in.pptx> <out.pptx> <slide-figure dir>"""
import os
import sys

from pptx import Presentation

sys.path.insert(0, os.path.dirname(__file__))
import deck_lib as L  # noqa: E402
from deck_lib import BODY, NAVY, RED  # noqa: E402

SRC, OUT, FIG = sys.argv[1], sys.argv[2], sys.argv[3]
prs = Presentation(SRC)
orig = list(prs.slides)
cover, tmpl, closing = orig[0], orig[1], orig[-1]
fig = lambda n: os.path.join(FIG, n)  # noqa: E731

X0, W0, Y0, H0 = 0.9, 11.53, 1.78, 4.62            # content area of the user's frame (inches)
LCOL, RCOL, CW = 0.9, 6.83, 5.6                   # two-column layout
new_slides = []


def content(title, foot=None):
    s = L.clone_slide(prs, tmpl)
    L.set_text(L.shape(s, "TextBox 2"), [{"k": "p", "t": title, "size": 22.5, "bold": True, "color": NAVY, "before": 0}])
    body = L.shape(s, "TextBox 4")
    L.place(body, X0, Y0, W0, H0)
    f = L.shape(s, "TextBox 5")
    if foot:
        L.set_text(f, [{"k": "n", "t": foot, "size": 11, "before": 0}])
    else:
        L.remove(f)
    new_slides.append(s)
    return s, body


def two_cols(s, body, left, right, lw=CW, rx=RCOL, rw=CW, y=Y0):
    L.place(body, LCOL, y, lw, H0)
    L.set_text(body, left)
    return L.textbox_like(s, body, rx, y, rw, H0, right)


def up(specs, d=1.0):
    """The same paragraphs ``d`` points larger (for slides with room to spare)."""
    out = []
    for sp in specs:
        sp = {"k": sp[0], "t": sp[1]} if isinstance(sp, tuple) else dict(sp)
        sp["size"] = sp.get("size", {"h": 16, "b": 14, "s": 13, "p": 14, "n": 12, "c": 14, "big": 26}[sp["k"]]) + d
        out.append(sp)
    return out


# ================================================================== cover (edited in place)
tb = L.shape(cover, "TextBox 2")
L.set_text(tb, [
    {"k": "c", "t": "TRƯỜNG CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG · ĐẠI HỌC BÁCH KHOA HÀ NỘI", "size": 13.5, "bold": True, "color": "FECACA", "before": 0},
    {"k": "c", "t": "BIỂU DIỄN NGỮ NGHĨA TỪ MÔ HÌNH NGÔN NGỮ LỚN", "size": 25, "bold": True, "color": "FFFFFF", "before": 6},
    {"k": "c", "t": "TRONG HỆ GỢI Ý TIN TỨC", "size": 25, "bold": True, "color": "FFFFFF", "before": 0},
    {"k": "c", "t": "Semantic Representations from LLMs for News Recommendation: A Controlled Study of 2024–2026 Directions on MIND",
     "size": 14, "italic": True, "color": "CBD5E1", "before": 6},
], lang="vi-VN")
L.notes(cover, "Kính chào Hội đồng. Đề tài nghiên cứu cách dùng biểu diễn ngữ nghĩa từ mô hình ngôn ngữ lớn cho gợi ý tin tức, "
        "đánh giá có kiểm soát các hướng LLM-rec 2024–2026 trên bộ dữ liệu MIND của Microsoft.")

# ================================================================== executive summary
s, body = content("Executive Summary")
cards = [("82–86%", "of test clicks hit articles never seen in training (MIND-small / large)"),
         ("0.4094 vs 0.4078", "best zero-shot LLM config vs retrained NRMS (MIND-small nDCG@10; CI includes 0)"),
         ("+0.0182", "category text in Qwen3-Embedding: the largest robust LLM-side gain"),
         ("0.038 → 0.007", "stacker validation optimism: user folds → time-block folds")]
for i, (big, lab) in enumerate(cards):
    L.card(s, X0 + i * 2.94, 1.82, 2.71, 1.5, big, lab)
L.place(body, X0, 3.5, W0, 2.9)
L.set_text(body, up([
    ("b", "**Frozen LLM embeddings match — but do not beat — properly retrained NRMS / NAML / Fastformer;** robust gains come from history aggregation (late interaction, +0.0128) and input fields (+0.0182), not from encoder size."),
    ("b", "**Most 2024–2026 LLM-rec directions we tested give ≤ 0.005 or negative effects** under strict evaluation: representation-level RAG / GraphRAG, generative semantic IDs, a 1.7B LLM judge, multi-interest models, LoRA, an LLM user tower."),
    ("b", "**A warm validation day misleads selection:** 18.1% cold clicks vs 82.2% at test; RAG / GraphRAG flip sign and learned readers are inflated."),
    ("b", "**MIND-large:** linear fusion 0.6863 AUC / 0.4373 nDCG@10; GBDT helps only with label-free exposure (0.7013 / 0.4516) and saturates; the ≈ 0.02 AUC gap to PLM-NR points to end-to-end fine-tuning → gated v17 run (pending)."),
], 1))
L.notes(s, "Bốn con số tóm tắt toàn bộ nghiên cứu. Thông điệp chính: embedding LLM đóng băng chỉ NGANG baseline huấn luyện đúng; lợi ích bền vững đến từ cách "
        "tổng hợp lịch sử và trường dữ liệu; phần lớn hướng LLM-rec 2024–2026 không giữ được lợi ích khi đánh giá nghiêm ngặt; khoảng trống còn lại nằm ở tinh chỉnh encoder.")

# ================================================================== 1. problem
s, body = content("1. Problem Formulation & Measured Challenges",
                  "* MINDsmall_dev: 73,152 and MINDlarge_dev: 376,471 test impressions; every number in this deck is measured on these sets.")
L.set_text(body, up([
    ("h", "Task: impression (slate) ranking [Wu et al., ACL'20]"),
    ("b", "User u with click history H = (h1 … hL); an impression shows candidates C = (c1 … cJ) with click labels y ∈ {0, 1}."),
    ("b", "Learn a score s(u, c); report AUC, MRR, nDCG@5 and nDCG@10 per impression, averaged (official MIND protocol)."),
    ("h", "Challenge 1 — item cold start dominates (measured)"),
    ("b", "85.9% (MIND-small) and 82.2% (MIND-large) of test clicks are on articles absent from every training row → ID-only collaborative filtering has no signal for most clicks; content representations are mandatory."),
    ("h", "Challenge 2 — chronological shift between validation and test (measured)"),
    ("b", "Validation = last training day: only 18.1% of its clicks are cold (MIND-large) → components that excel on seen articles look better than they are."),
    ("h", "Challenge 3 — short text, no article body"),
    ("b", "Titles 10.7 words, abstracts 34.9 words (5.2% empty); MIND ships URLs, not bodies → which fields reach the encoder matters (Section 4.3)."),
], 1))
L.notes(s, "Bài toán là xếp hạng các ứng viên trong một impression. Ba thách thức đều được đo trực tiếp từ dữ liệu: bài lạnh chiếm 82–86% lượt nhấp kiểm thử, "
        "validation ấm hơn test, và văn bản ngắn không có thân bài.")

# ================================================================== 2. RQs & contributions
s, body = content("2. Research Questions & Contributions")
two_cols(s, body, *[up(c, 0.5) for c in ([
    ("h", "Research questions"),
    ("b", "**RQ1** Where do zero-shot embeddings from modern encoders / LLMs stand against properly trained NRMS, NAML, Fastformer under one chronological protocol?"),
    ("b", "**RQ2** Is the bottleneck the encoder (size, family), the input fields, or history aggregation?"),
    ("b", "**RQ3** Do 2024–2026 LLM-rec directions (RAG / GraphRAG, semantic IDs, LLM user towers, LLM-generated knowledge, LLM judges, multi-interest, PEFT) survive strict evaluation on news?"),
    ("b", "**RQ4** At MIND-large scale: linear fusion or GBDT — and how trustworthy are validation estimates?"),
    ("b", "**RQ5** Where is the remaining gap to fine-tuned PLMs, and does partial fine-tuning close it? (v17, pending)"),
], [
    ("h", "Contributions"),
    ("b", "A controlled harness: chronological split, test touched once, paired 95% CIs per impression, cold / warm slices, measured seed noise, pre-registered predictions, 73 automated tests."),
    ("b", "≈ 16 technique families (≈ 150 variants) on MIND-small against retrained baselines; the V10 ↔ paper gap traced to pre-trained word vectors."),
    ("b", "Data findings: cold dominance, a warm validation day, the role of text fields, popularity as streaming feedback."),
    ("b", "Model findings: late interaction ≥ learned readers on small data; encoder size is not the bottleneck; GBDT saturates; user-fold CV is optimistic."),
    ("b", "A MIND-large pipeline on 1 × T4 / 32 GB with true peak-RAM accounting; honest positioning against published results."),
])])
L.notes(s, "Năm câu hỏi nghiên cứu, mỗi câu được trả lời bằng số đo ở các phần sau. RQ5 đang chờ kết quả thí nghiệm v17.")

# ================================================================== 3. landscape table
s, body = content("3. LLM-RecSys 2024–2026: Seven Directions We Examined",
                  "* Venues verified by web search (Oct 2026); the complete reference list is in the report and on the References slide.")
L.remove(body)
L.table(s, X0, 1.8, [1.95, 4.45, 2.9, 2.23], ["Direction", "Representative top-tier work", "Tested here", "Outcome on MIND"], [
    ["1. LLM as content encoder", "LLM2Vec (COLM'24); Qwen3-Embedding (2025); LLMEmb (AAAI'25); LLM2Rec (KDD'25); LEARN (AAAI'25); PAD (SIGIR'25)", "encoder sweep, concat encoder, learned readers (H8, H10)", "aggregation & fields > encoder size"],
    ["2. LLM-generated knowledge", "KAR (RecSys'24); ONCE (WSDM'24); EmbSum (RecSys'24); IRLLRec (SIGIR'25); PNR-LLM (WWW'25 Comp.)", "category / entity text (F); KAR zero-shot (H6); history-as-text (H1)", "+0.0182 (Qwen3); KAR mixed"],
    ["3. User modeling & multi-interest", "MINER (ACL-F'22); CAUM (SIGIR'22); Iana et al. (SIGIR'23); attention masks for LLM users (2026)", "LoRA user tower, 3 masks (H5); k-means & K-interest readers (H11)", "not reproducible / negative"],
    ["4. Generative rec & semantic IDs", "TIGER (NeurIPS'23); LETTER (CIKM'24); HSTU (ICML'24); ActionPiece (ICML'25); RPG (KDD'25); OneRec (2025); GRID (CIKM'25)", "RQ-KMeans 3×256 + SID-GPT slate ranker (H7)", "≈ random (0.2898 vs 0.2897)"],
    ["5. Reasoning", "ReaRec (2025); OneRec-Think (ACL'26); SIDReasoner (KDD'26); FLR, VRec, list-wise alignment (SIGIR'26)", "not tested (inference cost)", "—"],
    ["6. RAG / GraphRAG / agents", "GraphRAG (2024); K-RagRec (ACL'25); LLM-agent survey (EMNLP-F'25)", "representation RAG (H3); GraphRAG-lite + LLM summaries (H4)", "≤ +0.004; sign flips val → test"],
    ["7. LLM as ranker / judge", "LLMRank (ECIR'24); LLM4Rerank (WWW'25)", "Qwen3-1.7B Yes/No judge on top-10 (H9)", "−0.0112 (n = 2,000)"],
], size=11.5, row_h=0.52, bold_first=True, align=["l", "l", "l", "ctr"])
L.notes(s, "Bảng ánh xạ bảy hướng LLM-rec 2024–2026 (các hội nghị hàng đầu, đã xác minh venue) sang thí nghiệm tương ứng và kết quả đo trên MIND.")

# ================================================================== 3.1 directions 1-2
s, body = content("3.1 Directions 1–2: LLM Encoders & LLM-Generated Knowledge")
two_cols(s, body, [
    ("h", "Direction 1 — LLMs as content encoders"),
    ("b", "**LLM2Vec** (COLM'24): decoder-only LLM → encoder (bidirectional attention + contrastive training)."),
    ("b", "**Qwen3-Embedding** (2025): instruction-aware, last-token pooling, 0.6B–8B."),
    ("b", "**LLMEmb** (AAAI'25), **LLM2Rec** (KDD'25): adapt the LLM with supervised / collaborative fine-tuning, then use it as an item-embedding model."),
    ("b", "**LEARN** (AAAI'25): frozen LLM item encoder + twin-tower bridge to the collaborative space (industrial)."),
    ("b", "**PAD** (SIGIR'25): pre-train, align, disentangle text and collaborative views — largest gains on cold items."),
    ("b", "News: **Zhao et al.** (ECIR'25) — larger LMs are not reliably better on MIND."),
    ("n", "Tested here: frozen BGE-small/base/large, Qwen3-Emb-0.6B and their concatenation, zero-shot and as reader input."),
], [
    ("h", "Direction 2 — LLM-generated knowledge"),
    ("b", "**KAR** (RecSys'24): LLM reasoning + factual knowledge → hybrid-expert adaptor; deployed on Huawei news & music."),
    ("b", "**ONCE** (WSDM'24): open-source LLM as encoder + closed-source LLM for prompt augmentation (up to +19.32% rel.)."),
    ("b", "**EmbSum** (RecSys'24): LLM-supervised user-interest summaries → poly-embeddings."),
    ("b", "**Yada & Yamana** (2024): GPT-4 category descriptions on MIND (up to +5.8% AUC)."),
    ("b", "**PNR-LLM** (WWW'25 Comp.): hierarchical prompting to enrich news; **IRLLRec** (SIGIR'25): align LLM intents with interaction intents."),
    ("n", "Tested here: category / entity names in the embedded text (F); KAR-style item expansion and reader profiles with Qwen3-1.7B (H6)."),
])
L.notes(s, "Hướng 1: LLM làm bộ mã hoá — xu hướng 2025 là phải có bước thích nghi với hành vi (LLMEmb, LLM2Rec, PAD). Hướng 2: LLM sinh tri thức — KAR, ONCE, "
        "mô tả chuyên mục. Ghi chú màu xám là phần đã kiểm chứng trong nghiên cứu.")

# ================================================================== 3.2 directions 3-5
s, body = content("3.2 Directions 3–5: User Modeling, Generative Semantic IDs, Reasoning")
two_cols(s, body, [
    ("h", "Direction 3 — user modeling & multi-interest"),
    ("b", "Multi-interest networks: Li et al. (CIKM'19), ComiRec (KDD'20); news: **MINER** (ACL-F'22), **CAUM** (SIGIR'22)."),
    ("b", "**Iana et al.** (SIGIR'23): late fusion (dot products with clicked news) often beats complex user encoders."),
    ("b", "**Yuan et al.** (2026): attention masks for decoder-only user embeddings (causal / hybrid / bidirectional)."),
    ("n", "Tested here: LSE late interaction; k-means and K-interest readers; Qwen3 LoRA user tower with 3 masks."),
    ("h", "Direction 5 — reasoning (2025–2026)"),
    ("b", "**ReaRec** (2025) latent multi-step reasoning; **OneRec-Think** (ACL'26) interleaved text reasoning, deployed."),
    ("b", "**SIDReasoner** (KDD'26); SIGIR'26: **FLR**, **VRec** (reason–verify–recommend), Bayesian list-wise alignment."),
    ("n", "Not tested here: per-impression inference cost (conditional future work)."),
], [
    ("h", "Direction 4 — generative recommendation & semantic IDs"),
    ("b", "**TIGER** (NeurIPS'23): RQ-VAE semantic IDs + autoregressive retrieval; **LETTER** (CIKM'24): learnable tokenizer."),
    ("b", "**HSTU** (ICML'24): generative recommenders with compute scaling laws."),
    ("b", "2025: **ActionPiece** (ICML'25) context-aware tokenization; **RPG** (KDD'25) long semantic IDs in parallel; **OneRec** (Kuaishou) one generative model for retrieve + rank; **GRID** (CIKM'25) open framework."),
    ("b", "2026: SID-level evaluation can overstate item-level accuracy through code collisions (Zhang et al.)."),
    ("n", "Tested here: RQ-KMeans 3 × 256 codes + a 4-layer SID-GPT used as a slate ranker (H7)."),
])
L.notes(s, "Hướng 3 (mô hình hoá người dùng, đa sở thích), hướng 4 (gợi ý sinh với Semantic ID — dòng nóng nhất 2025–2026) và hướng 5 (suy luận). "
        "Hướng suy luận chưa thử do chi phí suy luận trên mỗi impression.")

# ================================================================== 3.3 directions 6-7, news PLMs, rigor
s, body = content("3.3 Directions 6–7, the News-PLM Line & Evaluation Rigor")
two_cols(s, body, *[up(c, 0.5) for c in ([
    ("h", "Direction 6 — RAG, GraphRAG, agents"),
    ("b", "**GraphRAG** (2024): community summaries over a knowledge graph; **K-RagRec** (ACL'25): KG retrieval for LLM recommenders."),
    ("b", "LLM-agent recommenders survey (EMNLP-F'25): recommender-, interaction- and simulation-oriented agents."),
    ("n", "Tested here: entity / kNN RAG on article vectors; Louvain communities + Qwen3-1.7B summaries."),
    ("h", "Direction 7 — LLM as ranker / judge"),
    ("b", "**LLMRank** (ECIR'24): zero-shot ranking with order and popularity biases; **LLM4Rerank** (WWW'25)."),
    ("n", "Tested here: Qwen3-1.7B Yes/No log-odds judge re-ranking the top-10."),
], [
    ("h", "News recommendation with PLMs"),
    ("b", "**PLM-NR** (SIGIR'21): PLM news encoder, fine-tune the last 2 layers; **UNBERT** (IJCAI'21); **GLORY** (RecSys'23)."),
    ("b", "Training cost is the bottleneck: **SpeedyFeed** (KDD'22, > 100× speed-up), **GRAM** (NAACL'22), **Tiny-NewsRec** (EMNLP'22)."),
    ("b", "**EB-NeRD** (RecSys Challenge'24) ships article bodies (MIND does not)."),
    ("h", "Evaluation rigor"),
    ("b", "Weak baselines inflate progress (**Dacrema et al.**, RecSys'19); a reproducibility study finds LLMRec behind established baselines (**Fioretti et al.**, RecSys'25)."),
    ("n", "→ Our protocol: retrained baselines, test touched once, paired CIs, measured seed noise."),
])])
L.notes(s, "Hướng 6–7 và dòng PLM cho tin tức. Bài học về độ tin cậy đánh giá (baseline yếu, không tái lập) là lý do giao thức của nghiên cứu nghiêm ngặt như vậy.")

# ================================================================== 4. data
s, body = content("4. Data: MIND-small vs MIND-large (measured from run logs)",
                  "* Counts printed by the data audit and loaders of runs v13 (small) and v14–v16b (large).")
L.table(s, X0, 1.8, [3.05, 1.85, 2.2], ["Statistic", "MIND-small", "MIND-large"], [
    ["Articles (train + dev merged)", "65,238", "104,151"],
    ["Official train impressions", "156,965", "2,232,748"],
    ["  train_core / validation (last day)", "126,695 / 30,270", "1,801,231 / 431,517"],
    ["Validation actually used", "30,270", "60k (v14); 150k (v15+)"],
    ["Test impressions (official dev)", "73,152", "376,471"],
    ["Training samples (click + 4 negatives)", "—", "2.71 M (0.6 M v14; 1.2 M v15+)"],
    ["Categories / subcategories", "18 / 270", "—"],
    ["Title / abstract length (words)", "10.7 / 34.9", "—"],
    ["Articles with entity names", "87.5% (2.4 / article)", "—"],
    ["Test readers' history length", "mean 32.3; 19.6% > 50", "—"],
    ["Test clicks on cold articles", "85.9%", "82.2% (val. 18.1%)"],
    ["Article body text", "not provided", "not provided"],
], size=11, row_h=0.33, highlight=(10,))
L.place(body, 8.35, 1.8, 4.08, H0)
L.set_text(body, [
    ("h", "Reading the table"),
    ("b", "Both test sets are overwhelmingly cold; the MIND-large validation day is not."),
    ("b", "Text is short and bodiless → field and aggregation choices matter."),
    ("b", "Most readers have < 50 clicks → longer histories barely help (+0.001)."),
    ("b", "MIND-large forced a lean loader, subsampled validation / training and a RAM-aware pipeline (Section 4.3)."),
])
L.notes(s, "Số liệu dữ liệu đo từ log các lần chạy. Điểm mấu chốt là dòng tô đỏ: tập kiểm thử rất lạnh, còn ngày validation của MIND-large thì ấm.")

# ================================================================== 4.1 cold
s, body = content("4.1 Finding 1: Test Clicks Are Cold, the Validation Day Is Warm")
L.picture(s, fig("s_cold.png"), X0, 1.95, 6.5, 3.9)
L.place(body, 7.65, Y0, 4.78, H0)
L.set_text(body, [
    ("h", "What we measured"),
    ("b", "Cold = the clicked article appears in no training row (history or impression)."),
    ("b", "Test: 85.9% (small) and 82.2% (large) of clicks are cold; MIND-large validation: 18.1%."),
    ("h", "Why it matters"),
    ("b", "On warm articles click-trained models win (warm AUC 0.56–0.64 vs 0.44–0.54 for zero-shot embeddings)."),
    ("b", "Frozen title + abstract embeddings have warm AUC < 0.5 (BGE-small 0.4416): similarity to the history anti-correlates with clicks on seen articles."),
    ("b", "Fix (v16): select fusion weights on a cold-weighted validation nDCG@10 (80% cold)."),
])
L.notes(s, "Khám phá dữ liệu thứ nhất: tập kiểm thử lạnh, validation ấm. Đây là gốc rễ của nhiều sai lệch khi chọn mô hình.")

# ================================================================== 4.2 val vs test
s, body = content("4.2 Finding 2: A Warm Validation Day Misleads Model Selection (v14)")
L.picture(s, fig("s_valtest.png"), X0, 1.85, 6.6, 4.4)
L.place(body, 7.75, Y0, 4.68, H0)
L.set_text(body, [
    ("h", "Gain over frozen: validation vs test"),
    ("b", "Zero-shot scorers keep most of their gain (test − validation: −0.005 … −0.009)."),
    ("b", "Learned readers lose 0.035–0.040 (um_knn_rag +0.065 → +0.030)."),
    ("b", "**kNN-RAG (+0.013 → −0.004) and GraphRAG (+0.020 → −0.007) flip sign.**"),
    ("h", "Consequence"),
    ("b", "Adding 3 readers + a 100-click history to the fusion raised validation (0.4152 → 0.4177) but lowered test: −0.0014 [−0.0017, −0.0011]."),
    ("b", "Every row is \"easier\" on test (even random vectors +0.020): compare deltas against a control, never absolute levels across splits."),
])
L.notes(s, "Khám phá thứ hai: lợi ích trên validation không chuyển sang test; RAG và GraphRAG còn đổi dấu. Vì vậy từ v16 chọn trọng số theo thước đo có trọng số bài lạnh.")

# ================================================================== 4.3 fields, popularity, memory
s, body = content("4.3 Findings 3–5: Text Fields, Popularity and Memory")
two_cols(s, body, [
    ("h", "Text fields matter (v13, MIND-small)"),
    ("b", "Dropping the abstract: −0.016 to −0.032 nDCG@10."),
    ("b", "Category / subcategory in the text: **Qwen3 +0.0182** [+0.0169, +0.0196], validation agrees; BGE-small +0.0013 (validation disagrees). Same on MIND-large: Qwen3 +0.0186, BGE-large −0.0030."),
    ("b", "The gain sits in the warm slice (warm AUC 0.4939 → 0.6180; cold +0.0005)."),
    ("b", "Entity names +0.0033; 100 / 200-click histories ≈ +0.001."),
    ("h", "Popularity: strong but streaming"),
    ("b", "Online popularity alone: 0.4026 (small) / 0.4040 (large) nDCG@10 — above every V10 neural model; it reads earlier clicks of the test day → kept in a separate \"+pop\" frame."),
], [
    ("h", "Scale & memory (MIND-large)"),
    ("b", "True peak RAM 21.2 GB (stacker stage; 32 GB container; GPU 5.0 GB), measured from v16b on."),
    ("b", "v15's \"16.2 GB\" was an end-of-stage snapshot, not a peak."),
    ("b", "v16 died after 1.6 h: the kNN-view gather allocated 2 × [104,151 × 20 × 2,048] float32 = 34.1 GB."),
    ("b", "Chunked computation (≈ 600 MB blocks): identical results (< 1e-6), +0.46 GB peak."),
    ("b", "Free Colab (~12.7 GB) is not enough at this scale; ≥ 24 GB is needed."),
])
L.notes(s, "Trường dữ liệu tạo khác biệt cỡ 0,02–0,03 — lớn hơn hầu hết kỹ thuật LLM phức tạp. Popularity mạnh nhưng là phản hồi streaming. "
        "Ở MIND-large, RAM là ràng buộc chính; lần chạy v16 chết vì phép gather 34 GB, đã sửa bằng tính theo khối.")

# ================================================================== 5. protocol
s, body = content("5. Evaluation Protocol & Leakage Controls")
L.set_text(body, up([
    ("h", "Chronological split, test touched once"),
    ("b", "Validation = last day of MIND_train; train_core = earlier days; MIND_dev = test. Every choice (τ, epochs, fusion weights, stacker configurations) uses validation only."),
    ("h", "Statistics"),
    ("b", "Paired 95% CI per impression: Δ̄ ± 1.96·SE with impression-clustered SE (analytic, not bootstrap); subset rows are compared on the same impressions; Bonferroni across technique families."),
    ("b", "Seed noise measured: 0.007–0.014 nDCG@10 between two seeds of one recipe — far wider than the CI half-width (± 0.002)."),
    ("h", "Reproduction checks against the original V10 benchmark"),
    ("b", "BGE-small mean-pool = V10 (AUC 0.6241 / nDCG@10 0.3909); popularity and frozen+pop match; NRMS rung 0 reproduces V10 (copied weights: |Δscore| ≤ 3·10⁻⁸); random vectors give AUC 0.50."),
    ("h", "Leakage tests (73 automated tests on CPU)"),
    ("b", "Shuffling test labels leaves stack_llm / stack_meta predictions bit-identical; flipping one fold's labels leaves that fold's out-of-fold scores unchanged; two deliberate leaks are caught. Predictions are written down before every run (Section 9.1)."),
], 1))
L.notes(s, "Giao thức: chia theo thời gian, test chạm một lần, CI ghép cặp theo impression (công thức giải tích có cụm, không phải bootstrap), đo nhiễu seed, "
        "đối chiếu tái lập với V10 và kiểm thử chống rò nhãn tự động.")

# ================================================================== 5.1 architecture
s, body = content("5.1 System Architecture",
                  "* Encoders are frozen in v7–v16b; only heads, readers and the stacker are trained. v17 unfreezes the last 2 layers of BGE-large.")
L.remove(body)
L.picture(s, fig("s_pipeline.png"), X0, 1.9, W0, 4.45)
L.notes(s, "Kiến trúc thí nghiệm: encoder đóng băng sinh vector bài; bốn họ thành phần; điểm được xuất và ghép tuyến tính hoặc bằng stacker GBDT; tất cả đánh giá dưới một giao thức.")

# ================================================================== 5.2 fusion & stacker
s, body = content("5.2 Linear Fusion and the GBDT Stacker")
two_cols(s, body, *[up(c, 0.5) for c in ([
    ("h", "Linear fusion (supermodel)"),
    ("b", "Each component exports raw scores; z-scored inside each impression."),
    ("b", "s_super = z_base + Σk wk·zk; base = best single component on validation (weight 1); wk ∈ {0, …, 4} by coordinate ascent; test touched once per rung."),
    ("b", "Ladder: super_base → super_zs (+ zero-shot) → super_llm (+ readers) → super_llm+pop."),
    ("b", "v16: base and weights chosen on cold-weighted validation nDCG@10 (80% cold)."),
    ("h", "Stacker variants"),
    ("b", "stack_llm = scores + ctx; stack_meta = + meta (label-free, reads impression-log exposure); stack_pop = + pop (streaming)."),
], [
    ("h", "GBDT stacker (XGBoost, GPU)"),
    ("b", "Trained on validation impressions only; test score = mean z-scored prediction of the fold models."),
    ("b", "**scores**: z-scores of every exported component (Qwen3, BGE-large, concat) + a few raw values."),
    ("b", "**ctx**: 11 history / slate features from frozen vectors (+ 4 IDF entity-overlap features in v16)."),
    ("b", "**meta** (16): item counts up to the previous day, smoothed CTR, category CTR, article age, prior impressions, 1/6/24-h exposure — no labels."),
    ("b", "**v15**: 4 user folds; objective chosen on fold 0 (binary:logistic); eta 0.06, depth 7, ≤ 800 trees."),
    ("b", "**v16**: time-block folds, selection on new-article impressions, 5 configurations on a forward holdout, blend of the best 2; the v15 recipe kept as a paired control."),
])])
L.notes(s, "Hai cách ghép: tuyến tính trên z-score và stacker XGBoost. Lưu ý mục tiêu huấn luyện thực tế là binary:logistic (chọn trên fold 0), không phải LambdaMART.")

# ================================================================== 6. catalogue
s, body = content("6. Techniques Tested: Catalogue and Verdicts")
L.remove(body)
L.table(s, X0, 1.78, [0.7, 2.2, 5.33, 3.3], ["Code", "Direction", "What we implemented", "Verdict (Δ nDCG@10 vs own control)"], [
    ["H1", "History as text", "10 latest titles → one instructed query, same encoder", "−0.0211 … −0.0014"],
    ["H2", "Co-click adaptation", "residual head, LoRA / DoRA, Matryoshka, impression hard negatives", "+0.0016 … +0.0029; cold AUC ↓"],
    ["H3", "Representation RAG", "article vector + Wikidata entities or k nearest past articles", "−0.0080 … −0.0003"],
    ["H4", "GraphRAG-lite", "Louvain communities; community summaries by Qwen3-1.7B", "+0.0043 (cold AUC −0.0031)"],
    ["H5", "LLM user tower", "Qwen3 + LoRA, InfoNCE; causal / bidirectional / soft masks", "mask effects not reproducible"],
    ["H6", "KAR-style knowledge", "Qwen3-1.7B item expansion / reader profiles (n = 2,500)", "+0.0079 / −0.0308"],
    ["H7", "Semantic IDs", "RQ-KMeans 3 × 256; prefix profile; SID-GPT slate ranker", "+0.0035; SID-GPT ≈ random"],
    ["H8", "Late interaction", "log-sum-exp, top-k, MaxSim, recency pooling", "+0.0128 (LSE)"],
    ["H9", "LLM judge", "Qwen3-1.7B Yes/No log-odds on the top-10 (n = 2,000)", "−0.0112 [−0.0216, −0.0008]"],
    ["H10", "Learned readers", "candidate-aware attention on frozen vectors; Fastformer / NRMS readers", "< LSE (small data); > LSE (large)"],
    ["H11", "Multi-interest", "k-means interests (K = 2, 3, 5); K-query readers ± diversity", "−0.0030 vs LSE; no gain"],
    ["F", "More data fields", "category, entity names, separate views, 100 / 200 clicks", "+0.0182 (Qwen3 category)"],
    ["R", "Retrained baselines", "NRMS / NAML / Fastformer, GloVe-300, published recipes, refit", "NRMS refit 0.4078"],
    ["S", "Fusion & stacking", "z-score linear fusion; XGBoost stacker (llm / meta / pop)", "Section 8"],
    ["FT", "Partial fine-tuning", "last 2 layers of BGE-large under the reader + frozen twin (v17)", "pending"],
], size=10.5, row_h=0.285, highlight=(7, 11), align=["ctr", "l", "l", "l"])
L.notes(s, "Danh mục mọi kỹ thuật đã thử cùng kết luận. Hai dòng tô đỏ là hai kết quả dương bền vững: late interaction và đưa chuyên mục vào văn bản.")

# ================================================================== 7. MIND-small main
s, body = content("7. MIND-small: Retrained Baselines vs Zero-shot LLM Embeddings",
                  "* Seed noise of a single trained model: 0.007–0.014 nDCG@10 (Section 7.2). Published row: mean of 5 runs.")
L.table(s, X0, 1.8, [3.05, 0.78, 0.78, 0.8, 0.85, 0.84], ["Model (content only)", "AUC", "MRR", "nDCG@5", "nDCG@10", "cold AUC"], [
    ["NRMS — V10 recipe (old)", "0.6181", "0.2809", "0.3020", "0.3685", "—"],
    ["NRMS — published recipe, GloVe", "0.6392", "0.3037", "0.3293", "0.3937", "0.6471"],
    ["NAML — published recipe", "0.6471", "0.3079", "0.3364", "0.4006", "0.6698"],
    ["Fastformer — mean of 2 seeds", "0.6426", "0.3000", "0.3289", "0.3927", "0.6623"],
    ["NRMS refit (+ validation days)", "0.6555", "0.3140", "0.3443", "0.4078", "0.6700"],
    ["BGE-small mean-pool (zero-shot)", "0.6242", "0.3002", "0.3299", "0.3909", "0.6617"],
    ["BGE-small + LSE", "0.6386", "0.3118", "0.3427", "0.4037", "0.6675"],
    ["Qwen3-Emb-0.6B + LSE + category", "0.6556", "0.3201", "0.3509", "0.4094", "0.6637"],
    ["Published NRMS (Iana et al., SIGIR'23)", "0.6658", "0.3144", "0.3499", "0.4121", "—"],
], size=11, row_h=0.4, highlight=(4, 7))
L.place(body, 8.4, Y0, 4.03, H0)
L.set_text(body, [
    ("h", "RQ1 — answer"),
    ("b", "Retrained baselines beat the old V10 numbers by up to +0.039 nDCG@10; NRMS refit sits 0.004 below the published value."),
    ("b", "Zero-shot BGE-small ≈ a one-seed baseline; late interaction lifts it to 0.4037."),
    ("b", "**Best zero-shot 0.4094 vs NRMS refit 0.4078: +0.0016, bounding CI [−0.0023, +0.0055] includes 0;** same AUC, +0.006 MRR / nDCG@5, lower cold AUC."),
    ("n", "Safe claim: frozen LLM embeddings match trained baselines without training or IDs — they do not beat them."),
])
L.notes(s, "Trả lời RQ1. Baseline huấn luyện đúng công thức mạnh hơn bộ đo V10 tới 0,039. Cấu hình zero-shot tốt nhất chỉ ngang NRMS huấn luyện lại, khoảng tin cậy chứa 0.")

# ================================================================== 7.1 what helps
s, body = content("7.1 MIND-small: What Actually Helps (and What Does Not)")
L.table(s, X0, 1.8, [4.05, 2.35, 1.05], ["Change (vs its own control)", "Δ nDCG@10 [95% CI]", "Val. agrees?"], [
    ["Late interaction LSE (BGE, vs mean-pool)", "+0.0128", "yes"],
    ["Late interaction top-k (BGE)", "+0.0107", "yes"],
    ["Category in text — Qwen3 (vs LSE)", "+0.0182 [+0.0169, +0.0196]", "yes"],
    ["Category in text — BGE-small (vs LSE)", "+0.0013 [+0.0000, +0.0026]", "no"],
    ["+ entity names — BGE (ctae vs cta)", "+0.0033 [+0.0025, +0.0041]", "yes"],
    ["200-click history — BGE (vs 50)", "+0.0012 [+0.0008, +0.0016]", "tiny"],
    ["Title only — Qwen3 (drop abstract)", "−0.0321 [−0.0338, −0.0304]", "yes"],
    ["Encoder: BGE-large vs BGE-small", "+0.0015 (n.s.)", "—"],
    ["Encoder: Qwen3-0.6B vs BGE-small", "−0.0125", "—"],
    ["Encoder: TF-IDF + LSA vs BGE-small", "−0.0443", "—"],
], size=11, row_h=0.4, highlight=(0, 2))
L.place(body, 8.65, Y0, 3.78, H0)
L.set_text(body, [
    ("h", "RQ2 — answer"),
    ("b", "The bottleneck of frozen-embedding recommenders is **history aggregation** and **which fields reach the encoder**, not encoder size."),
    ("h", "Late interaction (no parameters)"),
    ("p", "s_LSE(u, c) = τ·[log Σh exp(⟨vh, vc⟩ / τ) − log L]"),
    ("b", "τ → 0: MaxSim; τ → ∞: mean-pool; τ chosen on validation."),
    ("b", "Replicated on Qwen3: +0.0127."),
])
L.notes(s, "Trả lời RQ2: hai đòn bẩy rõ là late interaction (log-sum-exp, không tham số) và trường dữ liệu đưa vào encoder; encoder lớn hơn không giúp mean-pool zero-shot.")

# ================================================================== 7.2 ladder & seed noise
s, body = content("7.2 Why the Old Baselines Were Weak: Ablation Ladder & Seed Noise")
L.table(s, X0, 1.8, [6.25, 1.0, 1.1, 3.18], ["NRMS ablation rung (cumulative)", "AUC", "nDCG@10", "Δ nDCG@10 vs previous [95% CI]"], [
    ["Rung 0 — V10 recipe in the new harness", "0.6156", "0.3694", "—"],
    ["Rung 1 — + title-only, UNK, padding mask, negatives re-drawn every epoch", "0.6102", "0.3661", "−0.0033 [−0.0051, −0.0015]"],
    ["Rung 2 — + standard size / lr, random vectors (not converged)", "0.5818", "0.3447", "−0.0213 [−0.0229, −0.0197]"],
    ["NRMS, published recipe — + GloVe-300", "0.6392", "0.3937", "+0.0490 [+0.0471, +0.0509]"],
    ["NRMS refit — + training on the validation days", "0.6555", "0.4078", "+0.0141 [+0.0126, +0.0156]"],
], size=11, row_h=0.32, highlight=(3,))
L.table(s, X0, 4.1, [4.3, 1.0, 1.0, 1.0], ["Same recipe, two seeds", "seed A", "seed B", "gap"], [
    ["NRMS (published recipe)", "0.3937", "0.4005", "0.0068"],
    ["Qwen3 LoRA user tower (causal)", "0.3944", "0.3820", "0.0124"],
    ["Qwen3 reader um_frozen", "0.3929", "0.4063", "0.0134"],
    ["Qwen3 reader um_text", "0.4173", "0.4036", "0.0137"],
], size=11, row_h=0.3)
L.place(body, 8.55, 4.0, 3.88, 2.4)
L.set_text(body, [
    ("h", "Take-aways"),
    ("b", "Pre-trained word vectors (GloVe) explain most of the V10 ↔ paper gap; tokenizer / mask / negatives do not."),
    ("b", "A 'weak' baseline creates a 0.04 gap — larger than any LLM effect we measured."),
    ("b", "Seed noise ≫ paired CI: single-seed gaps < 0.007 are not findings."),
], first_before0=True)
L.notes(s, "Thang ablation cho thấy khoảng cách 0,04 giữa bộ đo V10 và bài báo đến từ việc thiếu vector GloVe. Nhiễu seed 0,007–0,014 lớn hơn nhiều so với CI ghép cặp.")

# ================================================================== 7.3 negatives
s, body = content("7.3 Null and Negative Results Across LLM Directions (MIND-small)")
L.picture(s, fig("s_directions.png"), X0, 1.75, 7.3, 4.75)
L.place(body, 8.4, Y0, 4.03, H0)
L.set_text(body, [
    ("h", "Negative or ≈ 0"),
    ("b", "Representation RAG −0.0080 … −0.0003; GraphRAG +0.0043 (cold AUC −0.0031)."),
    ("b", "SID-GPT slate ranker 0.2898 ≈ random vectors 0.2897."),
    ("b", "LLM-written reader profiles −0.0308; history-as-text −0.0211 … −0.0014."),
    ("b", "Qwen3-1.7B judge −0.0112 [−0.0216, −0.0008]; fused with the base score −0.0051 (n.s.)."),
    ("b", "LoRA / DoRA +0.0016 … +0.0029 with lower cold AUC."),
    ("h", "RQ3 — answer"),
    ("b", "Only the cheap form of LLM knowledge (category text) and late interaction give robust gains on news."),
])
L.notes(s, "Trả lời RQ3: biểu đồ cho thấy phần lớn các hướng LLM-rec nằm trong dải nhiễu hoặc âm. Semantic ID sinh tối giản ngang ngẫu nhiên; giám khảo LLM nhỏ có hại.")

# ================================================================== 7.4 multi-interest & tower
s, body = content("7.4 Multi-interest and the LLM User Tower: No Robust Gain")
L.table(s, X0, 1.8, [0.85, 3.1, 1.0, 0.8, 1.15], ["Enc.", "Reader (test nDCG@10, mean of seeds)", "nDCG@10", "std", "interest cos"], [
    ["BGE", "um_frozen (1 vector, 3 seeds)", "0.3968", "0.0016", "—"],
    ["BGE", "mi1 (K = 1 control)", "0.3906", "0.0035", "—"],
    ["BGE", "mi4 (K = 4)", "0.3944", "0.0011", "0.85–0.87"],
    ["BGE", "mi4d (K = 4 + diversity penalty)", "0.3972", "0.0055", "0.34–0.39"],
    ["Qwen3", "um_frozen (2 seeds)", "0.3996", "0.0067", "—"],
    ["Qwen3", "mi1 (K = 1 control)", "0.4008", "0.0018", "—"],
    ["Qwen3", "mi4 (K = 4)", "0.4029", "0.0019", "0.91"],
    ["Qwen3", "mi4d (K = 4 + diversity penalty)", "0.3965", "0.0032", "0.38–0.52"],
], size=11, row_h=0.3, align=["l", "l", "ctr", "ctr", "ctr"])
L.table(s, 8.2, 1.8, [2.3, 0.95, 0.98], ["Qwen3 + LoRA tower mask", "nDCG@10", "+ meanpool"], [
    ["causal (seed 0)", "0.3944", "0.4026"], ["bidirectional", "0.3854", "0.3965"], ["soft (v7)", "0.3964", "0.4043"],
    ["soft (v13 re-run)", "—", "0.3956"], ["causal (seed 1)", "0.3820", "0.3923"],
], size=11, row_h=0.3)
L.place(body, X0, 4.65, 6.9, 1.8)
L.set_text(body, [
    ("b", "k-means interests sit between mean-pool and LSE, monotone in K (−0.0079 / −0.0048 / −0.0030 vs LSE)."),
    ("b", "Without a penalty the K interests collapse (cos 0.85–0.91); with it they separate but ranking does not improve."),
])
L.textbox_like(s, body, 8.2, 3.75, 4.23, 2.6, [
    ("b", "soft − causal: +0.0020 (v7) → −0.0063 (v13), same code and seed."),
    ("b", "Seed gap 0.0124 ≥ any mask effect → no evidence that masks matter here (n = 30,000)."),
])
L.notes(s, "Đa sở thích và tower LLM: không có lợi ích bền vững. Khác biệt giữa các kiểu mặt nạ chú ý đổi dấu giữa hai lần chạy cùng mã — nhỏ hơn nhiễu seed.")

# ================================================================== 8. MIND-large v15
s, body = content("8. MIND-large (v15): Components, Linear Fusion and GBDT Stacker",
                  "* MINDlarge_dev, 376,471 impressions; stack_meta reads exposure in the impression log (no labels); the last three rows are streaming.")
L.table(s, X0, 1.8, [3.3, 0.82, 0.82, 0.88, 0.88, 0.88], ["Row", "AUC", "MRR", "nDCG@10", "cold AUC", "warm AUC"], [
    ["Qwen3 frozen (mean-pool)", "0.6215", "0.2922", "0.3787", "0.6583", "0.4784"],
    ["BGE-large frozen (mean-pool)", "0.6297", "0.3042", "0.3933", "0.6632", "0.4985"],
    ["super_llm (linear fusion)", "0.6774", "0.3286", "0.4268", "0.7004", "0.6207"],
    ["stack_llm (GBDT: scores + context)", "0.6843", "0.3355", "0.4332", "0.7003", "0.6534"],
    ["stack_meta (+ label-free exposure)", "0.6998", "0.3500", "0.4508", "0.7280", "0.6187"],
    ["frozen + pop (streaming)", "0.6757", "0.3324", "0.4312", "0.7259", "0.5098"],
    ["super_llm + pop (streaming)", "0.7114", "0.3574", "0.4599", "0.7446", "0.6192"],
    ["stack_pop (streaming)", "0.7206", "0.3680", "0.4730", "0.7524", "0.6313"],
], size=11, row_h=0.42, highlight=(3, 4))
L.place(body, 8.75, Y0, 3.68, H0)
L.set_text(body, [
    ("h", "Paired comparisons (nDCG@10)"),
    ("b", "stack_llm − super_llm: +0.0064 [+0.0059, +0.0069], but cold AUC −0.0001: the gain is on seen articles."),
    ("b", "stack_meta − stack_llm: +0.0176; cold AUC +0.0277, warm AUC −0.035."),
    ("b", "stack_pop − super_llm+pop: +0.0131."),
    ("h", "Caveat"),
    ("b", "Stacker out-of-fold validation sat 0.03–0.04 above test (0.4909 → 0.4508) — tested in v16b (Section 8.3)."),
])
L.notes(s, "Kết quả MIND-large v15. Stacker chỉ dùng nội dung hơn ghép tuyến tính 0,0064 nhưng không ở bài lạnh; bước nhảy lớn nhất đến từ đặc trưng phơi nhiễm không nhãn.")

# ================================================================== 8.1 readers by encoder
s, body = content("8.1 Readers by Encoder (v16b): A Bigger LLM Is Not Better")
L.picture(s, fig("s_readers.png"), X0, 1.85, 6.9, 4.45)
L.place(body, 8.1, Y0, 4.33, H0)
L.set_text(body, [
    ("h", "Same 128-d head, 2 seeds each"),
    ("b", "BGE-large (335M) − Qwen3-Emb (0.6B): **+0.0102** [+0.0095, +0.0108]."),
    ("b", "Concat − BGE-large: +0.0071 [+0.0066, +0.0076]; ordering holds for both seeds."),
    ("b", "kNN-RAG view: +0.0037 on BGE-large; +0.0003 on concat (n.s., cold AUC +0.0088)."),
    ("b", "With 1.2 M training samples a reader beats LSE on the same encoder (BGE-large +0.0105) — the reverse of MIND-small (−0.0080)."),
    ("n", "Diversity of representation spaces helps more than size (consistent with Zhao et al., ECIR'25)."),
])
L.notes(s, "Encoder LLM 0,6B kém encoder 335M khi làm đầu vào cho cùng bộ đọc; ghép hai encoder tốt nhất. Độ đa dạng biểu diễn có giá trị hơn kích thước.")

# ================================================================== 8.2 ladder
s, body = content("8.2 MIND-large Results Ladder (nDCG@10)")
L.picture(s, fig("s_ladder.png"), X0, 1.85, 7.45, 4.5)
L.place(body, 8.6, Y0, 3.83, H0)
L.set_text(body, [
    ("h", "Reading the ladder"),
    ("b", "Content + history ceiling ≈ 0.437: linear fusion and GBDT tie."),
    ("b", "Label-free exposure adds +0.0143 (stack_meta)."),
    ("b", "Streaming popularity adds +0.0215 more (stack_pop) — not comparable with offline papers."),
    ("b", "Colours mark the information each row uses, not the model family."),
])
L.notes(s, "Thang kết quả: trần của phần chỉ dùng nội dung và lịch sử khoảng 0,437; phơi nhiễm không nhãn thêm 0,014; popularity streaming thêm 0,022 nhưng không so được với số offline.")

# ================================================================== 8.3 stacker saturation & honest CV
s, body = content("8.3 Stacker Saturation and Honest Cross-Validation (v16b)")
L.picture(s, fig("s_stackcv.png"), X0, 1.85, 6.4, 4.4)
L.table(s, 7.55, 1.8, [2.55, 0.72, 0.85, 0.76], ["v16b row", "AUC", "nDCG@10", "v15"], [
    ["super_llm", "0.6863", "0.4373", "0.4268"], ["stack_llm", "0.6853", "0.4369", "0.4332"], ["stack_meta", "0.7013", "0.4516", "0.4508"],
    ["stack_meta_v15cfg (control)", "0.7029", "0.4543", "—"], ["stack_pop (streaming)", "0.7196", "0.4731", "0.4730"],
], size=11, row_h=0.3)
L.place(body, 7.55, 3.75, 4.88, 2.7)
L.set_text(body, [
    ("b", "**Saturation:** content improved +0.0105, stack_meta only +0.0008; stack_llm − super_llm = −0.0004 [−0.0010, +0.0002]."),
    ("b", "**Honest CV:** user folds overstate test by 0.038; time-block folds + new-article slice by 0.007."),
    ("b", "Yet the honest recipe is not better on test (−0.0027 [−0.0031, −0.0023] vs the v15 control): the *estimate* was wrong, not the model."),
])
L.notes(s, "Stacker đã bão hoà. Giả thuyết 'v15 tự đánh lừa mình' đúng về ước lượng (lệch 0,038), nhưng sửa nó không làm test tốt hơn: cái hỏng là ước lượng, không phải mô hình.")

# ================================================================== 8.4 gain shares
s, body = content("8.4 What the Trees Use: Split-Gain Shares")
L.picture(s, fig("s_gain.png"), X0, 1.85, 7.2, 4.4)
L.place(body, 8.4, Y0, 4.03, H0)
L.set_text(body, [
    ("h", "Interpretation"),
    ("b", "LLM component scores: 0.47–0.66 of total split gain in stack_llm / stack_meta (0.39–0.46 in stack_pop)."),
    ("b", "The context group (0.22–0.40) mixes LLM-derived similarities with non-LLM features — slate size (ctx_n_cand) alone takes 0.10–0.25."),
    ("b", "So the LLM share is not 'scores + context'; an '80% LLM' claim would overstate it."),
    ("b", "Top v16b features: cat:um_knn_rag (0.09–0.13), ctx_n_cand, bge:um_frozen, ctx_slate_max, m_ctr; streaming: pop, popw_ctr_1h."),
])
L.notes(s, "Phân tích gain: điểm LLM chiếm 47–66%. Không nên cộng nhóm ngữ cảnh vào 'phần của LLM' vì nhóm này có cả cỡ slate — một đặc trưng không liên quan LLM.")

# ================================================================== 9. positioning
s, body = content("9. Positioning Against Published MIND-large Results",
                  "* Published: leaderboard test set, secondary sources (PDFs not opened). Ours: MINDlarge_dev. ~0.01 AUC differences are not model claims.")
L.picture(s, fig("s_published.png"), X0, 1.72, W0, 3.3)
L.place(body, X0, 5.08, W0, 1.4)
L.set_text(body, [
    ("b", "Content + history rows (super_llm, stack_llm: AUC 0.684–0.686) lie inside the NRMS range reported in the literature (AUC 0.670–0.691) and ≈ 0.02 AUC below PLM-NR (0.7064) / UNBERT (0.7068)."),
    ("b", "stack_meta is within 0.005 AUC of PLM-NR and above it on nDCG@10 (0.4516 vs 0.4438), but reads impression-log exposure that PLM-NR / UNBERT do not use → a different information set; stack_pop needs test-day labels and cannot be submitted."),
])
L.notes(s, "Định vị trung thực: phần chỉ dùng nội dung thấp hơn PLM-NR khoảng 0,02 AUC. stack_meta dùng thêm thông tin phơi nhiễm nên không cùng bài toán; số công bố đo trên tập test khác.")

# ================================================================== 9.1 predictions
s, body = content("9.1 Pre-registered Predictions vs Outcomes",
                  "* Predictions were written in the README before each run; several of them turned out wrong, which is reported as is.")
L.remove(body)
L.table(s, X0, 1.8, [6.25, 3.1, 2.18], ["Prediction (written before the run)", "Measured", "Outcome"], [
    ["v13: category is the best new field, |Δ| ≲ 0.01", "Qwen3 +0.0182; BGE ≈ 0", "direction right, size wrong"],
    ["v13: k-means interests between mean-pool and LSE, possibly ≥ LSE", "between, always < LSE", "half right"],
    ["v15: stack_llm > super_llm, Δ in +0.003 … +0.02", "+0.0064", "right"],
    ["v15: stacker cold AUC > supermodel cold AUC", "−0.0001 (stack_llm)", "wrong"],
    ["v15: stack_meta > stack_llm (~60%)", "+0.0176", "right"],
    ["v16b: stack_meta AUC 0.698–0.712, nDCG@10 0.449–0.464", "0.7013; 0.4516", "inside"],
    ["v16b: stack_llm gains ≥ +0.01 over v15 (~15–20%)", "+0.0037", "wrong"],
    ["v16b: test − validation gap of stack_meta < control (~65%)", "−0.0072 vs −0.0379", "right"],
    ["v16b (implicit): stack_meta > stack_meta_v15cfg", "−0.0027", "wrong"],
    ["v17: PASS ~35% / INCONCLUSIVE ~25% / FAIL ~40%", "not run yet", "—"],
], size=12, row_h=0.4)
L.notes(s, "Mỗi phiên bản đều ghi dự đoán trước khi chạy, rồi đối chiếu với số đo. Việc công khai cả các dự đoán sai là một phần của tính nghiêm ngặt.")

# ================================================================== 10. compute
s, body = content("10. Compute, Memory and Reproducibility")
L.table(s, X0, 1.8, [0.75, 0.8, 4.2, 2.15], ["Run", "Data", "What changed", "Time / resources"], [
    ["V10", "small", "original benchmark: NRMS / NAML / Fastformer, BGE zero-shot, popularity", "—"],
    ["v7", "small", "≈ 16 families / ≈ 150 variants; retrained baselines; ablation ladder", "7.52 h (1 × T4)"],
    ["v13", "small", "text fields, 100 / 200 clicks, multi-interest, Fastformer × 2 + refit", "8.22 h"],
    ["v14", "large", "one Qwen3 model + all techniques; linear supermodel", "2.90 h"],
    ["v15", "large", "GBDT stacker (XGBoost GPU), exposure features, BGE-large", "2.73 h"],
    ["v16", "large", "time-block CV, BGE-large / concat readers", "died at 1.6 h (out of RAM)"],
    ["v16b", "large", "chunked kNN view, true peak RAM, 128-d heads for 3 encoders", "4.49 h; peak 21.2 GB"],
    ["v17", "large", "gated fine-tuning of BGE-large's last 2 layers", "est. 6–6.5 h (pending)"],
], size=11, row_h=0.42, align=["ctr", "ctr", "l", "l"])
L.place(body, 9.15, Y0, 3.28, H0)
L.set_text(body, [
    ("h", "Engineering notes"),
    ("b", "Each run is one 'Run All' Kaggle notebook (Tesla T4 16 GB, 32 GB container)."),
    ("b", "73 automated CPU tests before each run (correctness, leakage, memory)."),
    ("b", "Exported component scores + resume allow a stacker-only re-run without re-encoding the catalogue."),
    ("b", "Serving latency was **not** measured in this work."),
])
L.notes(s, "Mọi lần chạy là một notebook Kaggle (một GPU T4, RAM khoảng 32 GB); Colab miễn phí không đủ RAM cho MIND-large. "
        "Độ trễ phục vụ chưa được đo, nên không đưa ra con số độ trễ.")

# ================================================================== 11. answers & limitations
s, body = content("11. Answers to RQ1–RQ5 and Limitations")
two_cols(s, body, *[up(c, 1) for c in ([
    ("h", "Answers"),
    ("b", "**RQ1** Zero-shot ≈ retrained baselines (0.4094 vs 0.4078, n.s.)."),
    ("b", "**RQ2** History aggregation and input fields, not encoder size."),
    ("b", "**RQ3** Most directions ≤ 0.005 or negative; weak baselines and a warm validation day create false positives."),
    ("b", "**RQ4** Linear fusion ≈ GBDT on content; GBDT adds value only via label-free exposure and saturates ≈ 0.45; time-block CV is needed."),
    ("b", "**RQ5** The ≈ 0.02 AUC gap to PLM-NR points to end-to-end fine-tuning → v17 (pending)."),
], [
    ("h", "Limitations / threats to validity"),
    ("b", "One dataset family and one test day; dev instead of the leaderboard test; published numbers are secondary."),
    ("b", "Many rows single-seed; CIs exclude training noise (0.007–0.014)."),
    ("b", "No tuning beyond the stated grids; GloVe-6B; Fastformer re-implemented."),
    ("b", "LLM generation / judging only at 1.7B; Qwen3-Embedding 4B / 8B not tested; subset runs (n = 2,000–30,000)."),
    ("b", "The 80% cold target used knowledge of the test distribution (not labels)."),
])])
L.notes(s, "Tóm tắt câu trả lời năm câu hỏi nghiên cứu và các mối đe doạ tính hợp lệ.")

# ================================================================== 12. v17
s, body = content("12. Next Step: Gated Fine-tuning Experiment (v17)",
                  "* Gate and probabilities were written in the README before the run; only validation decides.")
two_cols(s, body, [
    ("h", "Design (PLM-NR recipe + SpeedyFeed-style caching)"),
    ("b", "Cache layer-22 hidden states of BGE-large for 104,151 articles ([104,151 × 96 × 1,024] fp16 ≈ 20 GB on local disk); check cos ≥ 0.99 against the full encoder."),
    ("b", "Train the last 2 layers + reader (fp16 autocast); candidate-side gradients; history table refreshed every 600 steps."),
    ("b", "Two matched arms from the same start: unfrozen (ft_reader) vs frozen twin (ft_base) — same steps, samples and order."),
    ("b", "Why the twin: in a toy run, continued reader training alone gave +0.0215 vs +0.0040 for unfreezing."),
], [
    ("h", "Decision rule (cold-weighted validation)"),
    ("b", "**PASS**: unfrozen − frozen ≥ +0.006 and raw mean not down > 0.001 → scale up (more layers / steps, both sides)."),
    ("b", "**INCONCLUSIVE**: ≥ +0.002 → add seeds first."),
    ("b", "**FAIL**: otherwise → keep frozen readers; write the thesis on v16b."),
    ("b", "If PASS (predicted): ft_reader 0.418–0.428 test nDCG@10 vs 0.4152 for the frozen BGE reader."),
])
for i, (big, lab, col) in enumerate([("PASS · ~35%", "≥ +0.006 → full fine-tuning run", "2E7D8C"),
                                     ("INCONCLUSIVE · ~25%", "≥ +0.002 → more seeds before deciding", "B7791F"),
                                     ("FAIL · ~40%", "keep frozen readers; thesis on v16b", RED)]):
    L.card(s, X0 + i * 3.9, 4.95, 3.67, 1.25, big, lab, big_color=col, big_size=21, label_size=12.5)
L.notes(s, "Bước tiếp theo là thí nghiệm có cổng dừng v17: tinh chỉnh hai lớp cuối của BGE-large với nhánh đối chứng đóng băng khớp cặp; "
        "quy tắc quyết định và xác suất dự đoán (PASS 35%, INCONCLUSIVE 25%, FAIL 40%) đã ghi trước khi chạy.")

# ================================================================== references
s, body = content("References (selected)", "* The complete, verified list (62 entries) is in the accompanying report.")
refs_l = ["[1] Wu et al. MIND: A Large-scale Dataset for News Recommendation. ACL 2020.",
          "[2] Wu et al. NRMS: Neural News Rec. with Multi-Head Self-Attention. EMNLP 2019.",
          "[3] Wu et al. NAML: Attentive Multi-View Learning. IJCAI 2019.",
          "[4] Wu et al. Fastformer: Additive Attention Can Be All You Need. arXiv 2021.",
          "[5] Wu et al. PLM-NR: Empowering News Rec. with PLMs. SIGIR 2021.",
          "[6] Zhang et al. UNBERT: User-News Matching BERT. IJCAI 2021.",
          "[7] Iana, Glavaš, Paulheim. Simplifying Content-Based Neural News Rec. SIGIR 2023.",
          "[8] Zhao, Huang, Vos, de Rijke. Revisiting LMs in Neural News Rec. ECIR 2025.",
          "[9] Zhang et al. Qwen3 Embedding. arXiv 2025.  [10] BehnamGhader et al. LLM2Vec. COLM 2024.",
          "[11] Liu et al. LLMEmb. AAAI 2025.  [12] He et al. LLM2Rec. KDD 2025.",
          "[13] Jia et al. LEARN. AAAI 2025.  [14] Wang et al. PAD. SIGIR 2025.",
          "[15] Xi et al. KAR. RecSys 2024.  [16] Liu et al. ONCE. WSDM 2024.",
          "[17] Yada, Yamana. Category Description by an LLM. arXiv 2024."]
refs_r = ["[18] Rajput et al. TIGER: Generative Retrieval. NeurIPS 2023.",
          "[19] Wang et al. LETTER: Learnable Item Tokenization. CIKM 2024.",
          "[20] Zhai et al. HSTU: Actions Speak Louder than Words. ICML 2024.",
          "[21] Hou et al. ActionPiece. ICML 2025.  [22] Hou et al. RPG. KDD 2025.",
          "[23] Deng et al. OneRec. arXiv 2025.  [24] Ju et al. GRID. CIKM 2025.",
          "[25] Liu et al. OneRec-Think. ACL 2026.  [26] He et al. SIDReasoner. KDD 2026.",
          "[27] Gao et al. Factorized Latent Reasoning. SIGIR 2026.  [28] Lin et al. VRec. SIGIR 2026.",
          "[29] Wang et al. K-RagRec. ACL 2025.  [30] Edge et al. GraphRAG. arXiv 2024.",
          "[31] Hou et al. LLMRank. ECIR 2024.  [32] Gao et al. LLM4Rerank. WWW 2025.",
          "[33] Xiao et al. SpeedyFeed. KDD 2022.  [34] Yu et al. Tiny-NewsRec. EMNLP 2022.",
          "[35] Ferrari Dacrema et al. Are We Really Making Much Progress? RecSys 2019.",
          "[36] Fioretti et al. A Reproducibility Study of LLMRec. RecSys 2025.",
          "[37] Kruse et al. EB-NeRD. RecSys Challenge 2024."]
two_cols(s, body, [{"k": "p", "t": r, "size": 12, "before": 3} for r in refs_l], [{"k": "p", "t": r, "size": 12, "before": 3} for r in refs_r])
L.notes(s, "Danh mục tài liệu chọn lọc; danh mục đầy đủ đã xác minh (62 mục) nằm trong báo cáo.")

# ================================================================== closing (edited in place) + drop the old content slides
tb = L.shape(closing, "TextBox 2")
L.set_text(tb, [
    {"k": "p", "t": "XIN TRÂN TRỌNG CẢM ƠN!", "size": 32, "bold": True, "color": NAVY, "before": 0},
    {"k": "p", "t": "Kính mong nhận được ý kiến trao đổi học thuật của Quý Thầy Cô và các bạn!", "size": 15.5, "italic": True, "color": RED, "before": 6},
    {"k": "p", "t": "TÓM LƯỢC BA LUẬN ĐIỂM KHOA HỌC CỐT LÕI:", "size": 15, "bold": True, "color": NAVY, "before": 14},
    {"k": "b", "t": "Data decide the evaluation: 82–86% of test clicks are cold and the validation day is warm — ignoring this selects the wrong models.", "size": 13.5, "before": 6},
    {"k": "b", "t": "Frozen LLM embeddings match, not beat, properly trained baselines; robust gains come from history aggregation and input fields, not encoder size.", "size": 13.5, "before": 6},
    {"k": "b", "t": "GBDT stacking adds value only through label-free exposure features and saturates (≈ 0.45 nDCG@10); the remaining gap is end-to-end fine-tuning (v17).", "size": 13.5, "before": 6},
    {"k": "p", "t": "Người thực hiện: NGUYỄN NAM HÁN   ·   MSHV: 20252095M", "size": 13.5, "bold": True, "color": NAVY, "before": 16},
    {"k": "p", "t": "Trường Công nghệ Thông tin và Truyền thông – ĐHBK Hà Nội", "size": 13.5, "color": BODY, "before": 2},
], lang="vi-VN")
L.notes(closing, "Ba luận điểm khoa học cốt lõi. Xin cảm ơn Hội đồng và mong nhận được góp ý.")

for sl in orig[1:-1]:
    L.delete_slide(prs, sl)
L.move_slide_to_end(prs, closing)
prs.save(OUT)
print("slides:", len(prs.slides), "->", OUT)
